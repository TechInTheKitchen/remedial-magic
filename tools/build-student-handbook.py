from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Handouts' / 'Prisms Student Handbook 5x7.docx'
OUT.parent.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(5), Inches(7)
sec.top_margin = sec.bottom_margin = Inches(.35)
sec.left_margin = sec.right_margin = Inches(.38)
sec.header_distance = sec.footer_distance = Inches(.15)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Subtitle'):
    s = doc.styles[name]
    s.font.name = 'Georgia' if name == 'Normal' else 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.04
doc.styles['Normal'].font.size = Pt(9.5)
doc.styles['Title'].font.size = Pt(20)
doc.styles['Title'].paragraph_format.space_after = Pt(3)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 2'].font.size = Pt(10.5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)
doc.core_properties.title = 'Prisms Student Handbook'
doc.core_properties.subject = 'Student record and player reference for Remedial Magic'
doc.core_properties.author = 'PRISMS Office of Summer Admissions'

def p(text='', bold=False, size=None, after=4):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(after)
    r = para.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    return para

def h(text): doc.add_paragraph(text, 'Heading 2')
def field(label, lines=1):
    para = p(label, bold=True, size=8.5, after=0)
    for _ in range(lines):
        para = p('________________________________________________________________', size=9, after=4)
        para.paragraph_format.line_spacing = 1.0
    return para

def table(headers, rows, widths, row_height=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    header = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(header)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row): cell.text = str(text)
    for i, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        if row_height and i:
            row.height = Inches(row_height)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e = OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+key),value)
                borders.append(e)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                m=OxmlElement('w:'+side); m.set(qn('w:w'),'55'); m.set(qn('w:type'),'dxa'); margins.append(m)
            props.append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9E9E9' if i==0 else ('F7F5F0' if i%2 else 'FFFFFF')); props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.0
                if j: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.name='Calibri'; r.font.size=Pt(8.5); r.bold=(i==0)
    p('', size=2, after=1)
    return t

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('PRISMS  |  STUDENT COPY  |  ').font.size = Pt(7)
page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'),'PAGE'); footer._p.append(page)

# FRONT COVER: all core student records remain on this page.
p('PENGORGON\u2019S REMEDIAL INSTITUTE', bold=True, size=8, after=1)
p('FOR SUMMER MAGICAL STUDIES', size=8, after=5)
doc.add_paragraph('Prisms Student Handbook', 'Title')
p('Your record, component allowance, and instructions for unscheduled practical work.', size=9, after=7)
field('Student name and preferred form of address')
field('Admitted following', 2)
field('Specialty school                           Favored color')
field('Background knack')
field('Outstanding business   Who or what still needs you', 2)
table(['Level', 'Ordinary draw', 'Specialty draw'], [['1 / ___', '5', '7 / ___']], [1.15,1.45,1.64], .25)
p('COMPOSURE   10 maximum', bold=True, size=9, after=3)
p('\u25cb  \u25cb  \u25cb  \u25cb  \u25cb     \u25cb  \u25cb  \u25cb  \u25cb  \u25cb', size=15, after=5)
p('Mark lost Composure in pencil. At zero, you are Dazed.', size=8, after=4)
field('Current pouch total                 No color above half')
field('Current enchantment or impediment')
field('Identifying features currently present or missing')

doc.add_page_break()
doc.add_paragraph('Your Component Allowance', 'Heading 1')
p('One pouch. All supplies mixed together. Record color counts, not individual dice.', after=7)
field('Expedition                           Preferred colors')
table(['Color','Own','Loan','Found','Extra','Total'], [
    ['Red','','','','',''],['Blue','','','','',''],['Green','','','','',''],
    ['Brown','','','','',''],['White','','','','',''],['Purple','','','','',''],
    ['Black','','','','',''],['Totals','','','','','']], [1.04,.60,.64,.66,.66,.64], .31)
p('Own = permanent. Loan = practical packet. Extra = corrective supplies. Every row adds up to Total.', size=8, after=5)
h('Collecting your allowance')
p('Start with 12 dice: 4 school-favored, 4 of a different chosen color, and 4 random. Allocate random colors using the table below; reroll 8 or any illegal allocation. No color may exceed half the completed pouch, rounded down.')
p('D8 allocation: 1 Red / 2 Blue / 3 Green / 4 Brown / 5 White / 6 Purple / 7 Black.', size=8.5)
h('Expedition supplies')
p('At Levels 1 to 5, mix in 6 / 8 / 10 / 12 / 14 loan dice, randomly chosen outside your preferred colors. These supplies are compulsory, even when your destination is not approved.')
p('Found   \u25a1 \u25a1 \u25a1     Corrective   \u25a1 \u25a1 \u25a1 \u25a1', bold=True, size=10)
p('Find up to 1 ingredient per significant scene, 3 per expedition. The Professor names its color before you accept. Getting caught may bring 2 random unwanted dice, up to 4 corrective dice per expedition. Respect the half-pouch limit; additions never enlarge your draw.', size=9)

doc.add_page_break()
doc.add_paragraph('Approved Casting Procedure', 'Heading 1')
p('Quill need not approve. The Professor running the game must know what you are attempting.', size=9)
h('Declare before drawing')
p('State effect, target, school, and intended mana. Agree difficulty. Draw 5 blindly, or your specialty allowance. Roll the full draw without selecting, redrawing, or rerolling dice.')
table(['Difficulty','Matching faces','Strength'], [['Routine','2 alike','2'],['Demanding','3 alike','4'],['Exceptional','4 alike','6']], [1.54,1.65,1.05])
p('Any face value qualifies. Use the largest group; two pairs are not a triple. Exceed the required group for +1 strength, once. Agree scope before drawing: usually 1 / up to 3 / up to 5 targets as difficulty rises.', size=9)
h('Count every color')
p('Intended color leads alone: your intended spell works. Another color leads alone: a useful unexpected substitution works. Tied leading colors: Wild magic gives a useful effect and a complication. Faces below difficulty: one backfire, themed by the leading color or tied colors.')
p('The caster chooses the surprise, within the colors and permitted scope. Ask classmates for ideas if needed; the Professor chooses when you need help. An obvious twist already present in the scene may supply the explanation. Success stays useful.')
h('Record one principal effect')
p('Strength sets attack damage, healing, or shield protection. Divide points among declared targets before drawing. An extra point goes to an existing target. Utility magic follows scope. Shields do not stack; use the strongest remaining shield.')
p('Shortfall of 1 / 2 / 3 dice: brief / scene-lasting / major trouble, OR 1 / 2 / 3 damage. No automatic double punishment. Missing eyebrows are complimentary.', size=9)
p('RETURN EVERY DIE AND MIX AFTER EVERY CAST', bold=True, size=9)
p('Red fire / Blue water and binding / Green life / Brown earth / White air and wards / Purple mind / Black spirits.', size=8.5)

doc.add_page_break()
doc.add_paragraph('Conduct During an Incident', 'Heading 1')
p('Quill naps. You slip out to help friends or investigate. Your knack establishes safe ordinary competence; risky tasks use a charm and the normal casting procedure. An open door needs no spell.')
h('When orderly turns are required')
p('Agree student order. Each gets 1 meaningful action, ordinary movement, and brief speech; then the announced threat acts. Major finales may have 2 announced actions. Describe reach and obstacles rather than measuring distance.')
p('Help: spend your action naming a classmate and task. Their next relevant draw gains 1 die from their own pouch. No helper roll or stacking; expires at the end of the next student phase.')
h('Nurse Bell and the attention register')
p('At 0 Composure, speak but take no normal action. Collection follows at the end of the next student phase after you fall. A safely reachable classmate can spend an action restoring 1 Composure without a roll; healing also cancels collection if you rise above zero.')
p('Collection restores all 10 and returns you at a sensible scene transition. A safe breather after a pressured scene restores 2 once; proper rest restores all 10. Excess shield damage reduces Composure.')
p('Suspicion is shared: 0\u20131 unnoticed / 2\u20133 inquiries / 4\u20135 investigation / 6 intervention, one consequence, then reset to 3. Clear evidence adds 1 per incident, not every backfire. Debrief resets it to zero.', size=8.5)
h('Debrief and promotion')
p('Return all recorded loan, found, and corrective dice. Major practicals advance the cohort, up to Level 5, adding 2 chosen permanent dice per level gained. Specialty draws: 7 at Levels 1\u20132, 8 at 3\u20134, 9 at 5. Permanent totals: 12 / 14 / 16 / 18 / 20. At Level 5, return loans; later debriefs may recolor 2 existing dice. Keep the half-pouch limit.', size=8.5)
field('What I intended and what actually happened')
field('Faculty remarks or next outstanding business')
p('Student signature __________________________', size=9, after=0)
p('Signing confirms attendance, not agreement with the shrub.', size=8, after=0)

doc.save(OUT)
print(OUT)
